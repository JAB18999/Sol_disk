#!/usr/bin/env python3
"""Download the latest confirmed OKX SOL-USDT-SWAP candles.

Creates four CSV files containing exactly 8,640 completed candles each:
15m, 30m, 1H, and 2H.  The script keeps the API's raw OHLCV strings in
CSV and writes a metadata file with coverage and gap checks.
"""

from __future__ import annotations

import csv
import hashlib
import json
import time
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from zoneinfo import ZoneInfo

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
except ImportError:
    Workbook = None

ROOT = Path(__file__).resolve().parent
INSTRUMENT = "SOL-USDT-SWAP"
COUNT = 8640
BARS = {
    "15m": 15 * 60 * 1000,
    "30m": 30 * 60 * 1000,
    "1H": 60 * 60 * 1000,
    "2H": 2 * 60 * 60 * 1000,
}
API_BASE = "https://www.okx.com/api/v5/market/history-candles"
INSTRUMENT_URL = "https://www.okx.com/api/v5/public/instruments"
HEADERS = {
    "User-Agent": "Mozilla/5.0 (compatible; Arena-OKX-Kline-Downloader/1.0)",
    "Accept": "application/json",
}
CSV_HEADER = [
    "open_time_utc",
    "open_time_shanghai",
    "timestamp_ms",
    "open",
    "high",
    "low",
    "close",
    "vol",
    "volCcy",
    "volCcyQuote",
    "confirm",
]


def fetch_json(url: str, attempts: int = 6) -> dict:
    """Fetch one public OKX endpoint with bounded retries."""
    last_error = None
    for attempt in range(attempts):
        request = urllib.request.Request(url, headers=HEADERS)
        try:
            with urllib.request.urlopen(request, timeout=30) as response:
                payload = json.loads(response.read().decode("utf-8"))
            if payload.get("code") != "0":
                raise RuntimeError(f"OKX returned code={payload.get('code')}: {payload.get('msg')}")
            return payload
        except (urllib.error.URLError, urllib.error.HTTPError, TimeoutError, RuntimeError, json.JSONDecodeError) as exc:
            last_error = exc
            # Mild backoff also avoids public-endpoint rate limits.
            time.sleep(min(8, 0.75 * (2 ** attempt)))
    raise RuntimeError(f"Unable to fetch OKX after {attempts} attempts: {last_error}")


def iso_times(timestamp_ms: int) -> tuple[str, str]:
    utc = datetime.fromtimestamp(timestamp_ms / 1000, tz=timezone.utc)
    shanghai = utc.astimezone(ZoneInfo("Asia/Shanghai"))
    return utc.isoformat().replace("+00:00", "Z"), shanghai.isoformat()


def get_candles(bar: str, required_count: int) -> tuple[list[list[str]], int]:
    """Return newest `required_count` confirmed candles, sorted oldest to newest."""
    confirmed_newest_first: list[list[str]] = []
    seen_timestamps: set[int] = set()
    cursor: int | None = None
    request_count = 0

    while len(confirmed_newest_first) < required_count:
        params = {"instId": INSTRUMENT, "bar": bar, "limit": "300"}
        if cursor is not None:
            # For history-candles, `after` returns records strictly earlier than this timestamp.
            params["after"] = str(cursor)
        url = f"{API_BASE}?{urllib.parse.urlencode(params)}"
        payload = fetch_json(url)
        page = payload.get("data", [])
        request_count += 1
        if not page:
            raise RuntimeError(f"OKX returned no more {bar} data after cursor {cursor}; only {len(confirmed_newest_first)} completed candles collected.")

        page_timestamps = [int(candle[0]) for candle in page]
        if cursor is not None and max(page_timestamps) >= cursor:
            raise RuntimeError(f"Unexpected non-decreasing pagination for {bar}: cursor={cursor}, page_max={max(page_timestamps)}")

        # The endpoint returns newest-to-oldest. Keep that ordering until the latest N are chosen.
        for candle in page:
            timestamp_ms = int(candle[0])
            if candle[8] == "1" and timestamp_ms not in seen_timestamps:
                confirmed_newest_first.append(candle)
                seen_timestamps.add(timestamp_ms)

        cursor = min(page_timestamps)
        # Public endpoint limit is 20 requests / 2 seconds.  This keeps below it.
        time.sleep(0.115)

    latest_confirmed = confirmed_newest_first[:required_count]
    if len({int(c[0]) for c in latest_confirmed}) != required_count:
        raise RuntimeError(f"Duplicate timestamps found in selected {bar} data")

    return sorted(latest_confirmed, key=lambda c: int(c[0])), request_count


def csv_rows(candles: list[list[str]]) -> list[list[str]]:
    result = []
    for candle in candles:
        ts, o, h, l, c, vol, vol_ccy, vol_ccy_quote, confirm = candle
        utc, shanghai = iso_times(int(ts))
        result.append([utc, shanghai, ts, o, h, l, c, vol, vol_ccy, vol_ccy_quote, confirm])
    return result


def check_integrity(candles: list[list[str]], expected_step_ms: int) -> dict:
    timestamps = [int(c[0]) for c in candles]
    steps = [later - earlier for earlier, later in zip(timestamps, timestamps[1:])]
    anomalies = [
        {
            "previous_timestamp_ms": timestamps[i],
            "current_timestamp_ms": timestamps[i + 1],
            "observed_step_ms": step,
            "expected_step_ms": expected_step_ms,
        }
        for i, step in enumerate(steps)
        if step != expected_step_ms
    ]
    return {
        "unique_timestamp_count": len(set(timestamps)),
        "expected_step_ms": expected_step_ms,
        "interval_anomaly_count": len(anomalies),
        "interval_anomalies": anomalies[:20],
        "all_confirmed": all(c[8] == "1" for c in candles),
    }


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for block in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def write_workbook(all_rows: dict[str, list[list[str]]], metadata: dict) -> Path | None:
    if Workbook is None:
        return None

    output = ROOT / f"{INSTRUMENT}_OKX_8640bars_4timeframes.xlsx"
    workbook = Workbook(write_only=False)
    summary = workbook.active
    summary.title = "说明"
    summary.append(["OKX SOL 永续合约 K 线数据"])
    summary.append(["合约", INSTRUMENT])
    summary.append(["数据源", API_BASE])
    summary.append(["下载时间（UTC）", metadata["downloaded_at_utc"]])
    summary.append(["筛选条件", "仅 confirm=1 的已收盘 K 线；每个周期 8,640 行；时间升序"])
    summary.append([])
    summary.append(["周期", "行数", "开始时间（UTC）", "结束时间（UTC）", "间隔异常数"])
    for bar in BARS:
        info = metadata["timeframes"][bar]
        summary.append([
            bar,
            info["row_count"],
            info["start_time_utc"],
            info["end_time_utc"],
            info["integrity"]["interval_anomaly_count"],
        ])
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 25
    summary.column_dimensions["C"].width = 42
    summary.column_dimensions["D"].width = 42
    summary.freeze_panes = "A7"

    for bar, rows in all_rows.items():
        sheet = workbook.create_sheet(bar)
        sheet.append(CSV_HEADER)
        for row in rows:
            # Keep timestamps as text; numeric prices/volumes are written as numbers for spreadsheet use.
            formatted = row[:3] + [float(value) for value in row[3:10]] + [int(row[10])]
            sheet.append(formatted)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:K{len(rows) + 1}"
        for col, width in {
            "A": 23, "B": 30, "C": 16, "D": 13, "E": 13, "F": 13,
            "G": 13, "H": 16, "I": 16, "J": 18, "K": 10,
        }.items():
            sheet.column_dimensions[col].width = width
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")

    for cell in summary[1]:
        cell.font = Font(bold=True, size=14)
    workbook.save(output)
    return output


def main() -> None:
    ROOT.mkdir(parents=True, exist_ok=True)
    downloaded_at = datetime.now(timezone.utc)
    metadata: dict = {
        "instrument": INSTRUMENT,
        "instrument_type": "USDT-margined perpetual swap",
        "source_endpoint": API_BASE,
        "source_parameters": {"instId": INSTRUMENT, "limit_per_request": 300},
        "downloaded_at_utc": downloaded_at.isoformat().replace("+00:00", "Z"),
        "downloaded_at_shanghai": downloaded_at.astimezone(ZoneInfo("Asia/Shanghai")).isoformat(),
        "completed_candles_only": True,
        "sort_order": "ascending by candle open timestamp",
        "requested_rows_per_timeframe": COUNT,
        "columns": CSV_HEADER,
        "timeframes": {},
    }

    # Record the public instrument specification when available.
    instrument_params = urllib.parse.urlencode({"instType": "SWAP", "instId": INSTRUMENT})
    instrument_payload = fetch_json(f"{INSTRUMENT_URL}?{instrument_params}")
    metadata["instrument_snapshot"] = instrument_payload.get("data", [])

    all_rows: dict[str, list[list[str]]] = {}
    csv_paths: list[Path] = []
    for bar, bar_ms in BARS.items():
        print(f"Downloading {bar}: latest {COUNT} confirmed candles ...", flush=True)
        candles, request_count = get_candles(bar, COUNT)
        integrity = check_integrity(candles, bar_ms)
        if integrity["unique_timestamp_count"] != COUNT or not integrity["all_confirmed"]:
            raise RuntimeError(f"Integrity check failed for {bar}: {integrity}")
        if integrity["interval_anomaly_count"]:
            raise RuntimeError(f"Found interval gaps/anomalies for {bar}: {integrity['interval_anomalies'][:3]}")

        rows = csv_rows(candles)
        all_rows[bar] = rows
        csv_path = ROOT / f"{INSTRUMENT}_{bar}_{COUNT}_confirmed.csv"
        with csv_path.open("w", newline="", encoding="utf-8") as out:
            writer = csv.writer(out)
            writer.writerow(CSV_HEADER)
            writer.writerows(rows)
        csv_paths.append(csv_path)

        metadata["timeframes"][bar] = {
            "bar": bar,
            "row_count": len(candles),
            "api_request_count": request_count,
            "start_timestamp_ms": candles[0][0],
            "end_timestamp_ms": candles[-1][0],
            "start_time_utc": rows[0][0],
            "end_time_utc": rows[-1][0],
            "start_time_shanghai": rows[0][1],
            "end_time_shanghai": rows[-1][1],
            "integrity": integrity,
            "csv_filename": csv_path.name,
            "csv_sha256": sha256(csv_path),
        }
        print(f"  saved {csv_path.name}; {rows[0][0]} -> {rows[-1][0]}", flush=True)

    metadata_path = ROOT / "metadata.json"
    metadata_path.write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")

    readme = ROOT / "README.md"
    lines = [
        "# OKX SOL-USDT-SWAP K 线数据",
        "",
        "- 数据源：OKX Public API `GET /api/v5/market/history-candles`",
        f"- 合约：`{INSTRUMENT}`（USDT 本位永续合约）",
        f"- 下载时间（UTC）：{metadata['downloaded_at_utc']}",
        "- 周期：15m、30m、1H、2H；每个 CSV 均为最新的 8,640 根已收盘 K 线。",
        "- 已剔除 API 返回中 `confirm=0` 的当前未收盘 K 线。",
        "- 所有文件按 K 线开盘时间升序排列，时间戳表示 K 线开盘时刻。",
        "- 完整性检查：每个周期 8,640 个唯一时间戳，连续间隔无缺口。",
        "",
        "## CSV 字段",
        "",
        "`open_time_utc`、`open_time_shanghai`、`timestamp_ms`、`open`、`high`、`low`、`close`、`vol`、`volCcy`、`volCcyQuote`、`confirm`。",
        "",
        "其中 `vol`、`volCcy`、`volCcyQuote` 保留 OKX K 线 API 的原始字段名称和数值；`confirm=1` 表示 K 线已收盘。",
        "",
        "## 覆盖区间",
        "",
        "| 周期 | 行数 | 开始（UTC） | 结束（UTC） |",
        "|---|---:|---|---|",
    ]
    for bar in BARS:
        info = metadata["timeframes"][bar]
        lines.append(f"| {bar} | {info['row_count']:,} | {info['start_time_utc']} | {info['end_time_utc']} |")
    lines.extend([
        "",
        "原始 CSV 用于程序化分析；同目录的 `.xlsx` 工作簿将四个周期分别放在四个工作表中。",
        "下载脚本 `download_okx_sol_perp_klines.py` 一并保留，便于以后复现下载过程。",
    ])
    readme.write_text("\n".join(lines) + "\n", encoding="utf-8")

    workbook_path = write_workbook(all_rows, metadata)

    # Bundle raw CSV data, metadata, documentation, and the reproducible downloader.
    zip_path = ROOT / f"{INSTRUMENT}_OKX_8640bars_4timeframes_raw_csv.zip"
    with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as archive:
        for path in csv_paths + [metadata_path, readme, Path(__file__)]:
            archive.write(path, arcname=path.name)

    print("\nCompleted successfully.")
    for bar in BARS:
        info = metadata["timeframes"][bar]
        print(f"{bar}: {info['row_count']} rows, {info['start_time_utc']} -> {info['end_time_utc']}")
    print(f"Raw-data ZIP: {zip_path.name}")
    if workbook_path:
        print(f"Workbook: {workbook_path.name}")


if __name__ == "__main__":
    main()
